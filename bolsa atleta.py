import pandas as pd

from api_consultas.funtions_main import *
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
import locale
from datetime import datetime, timedelta
from docx2pdf import convert
from cache_manager import *

locale.setlocale(locale.LC_TIME, 'pt_BR.utf8')


def create_declaration_prestacao(tabela_prestacao_ministerio):

    df_atleta = CACHE(cache_file_name='main_atlleta').load_dataframe_from_cache()
    df_cpf = CACHE(cache_file_name='cpf_id_atlleta').load_dataframe_from_cache()

    atleta_cpf_df = pd.merge(df_atleta,
                             df_cpf,
                             how='left',
                             left_on='id',
                             right_on='id_atleta')

    results_sge = CACHE(cache_file_name='todos_dados_api_rank_arena').load_dataframe_from_cache()
    eventos_df = CACHE(cache_file_name='dados_eventos_2023a2025').load_dataframe_from_cache()

    results_sge = results_sge.merge(eventos_df, how='left', left_on='id_evento', right_on='id')

    tabela_sge_ministerio = pd.merge(tabela_prestacao_ministerio,
                                     atleta_cpf_df,
                                     how='left',
                                     left_on='cpf_normal',
                                     right_on='cpf')

    print(tabela_sge_ministerio)

    for index, atleta in tabela_sge_ministerio.iterrows():

        doc = Document(
            r"C:\Users\agata\CBW 2024\Bolsa Atleta\MODELO DECLARAÇÃO DA CONFEDERAÇÃO BRASILEIRA DE WRESTLING - PRESTAÇÃO DE CONTAS 2023 com placeholders.docx")

        nome_do_cidadao = atleta['nome_completo']
        id_do_cidadao = atleta['id_atleta']

        for paragraph in doc.paragraphs:

            print("paragáfo:", paragraph.text)

            if "{{hash1}}" in paragraph.text:

                texto_base = f"Declara para fins de PRESTAÇÃO DE CONTAS, que o (a) atleta {nome_do_cidadao}, inscrito(a) sob o CPF nº {atleta['CPF']}, beneficiado com a Bolsa Atleta na categoria {atleta['Categoria de Bolsa']}."
                paragraph.text = texto_base
                run = paragraph.runs[0]
                font = run.font
                font.name = 'Calibri'
                font.size = Pt(11)
                font.bold = False

                print("paragáfo alterado:", paragraph.text)

            elif "{{hash2}}" in paragraph.text:

                lista_eventos = []

                df_loc = results_sge[results_sge['fullName'] == nome_do_cidadao]

                print(df_loc)

                try:

                    for index, nome in df_loc.iterrows():

                        nome['data_inicio'] = pd.to_datetime(nome['data_inicio'], dayfirst=True)
                        atleta['Fim do Pagamento'] = pd.to_datetime(atleta['Fim do Pagamento'], dayfirst=True)
                        atleta['Data início pagamento'] = atleta['Inicio do Pagamento'] # + pd.DateOffset(years=-1)

                        print("temos resultados")

                        if (nome["fullName"] == nome_do_cidadao and atleta['Data início pagamento'] <= nome['data_inicio'] <= atleta['Fim do Pagamento']):
                            data_evento = datetime.strptime(str(nome['data_fim']), "%Y-%m-%d %H:%M:%S").strftime(
                                "%d de %B de %Y")
                            texto_eventos = f"• {nome['descricao']}, {int(nome['rank'])}º colocado(a) na categoria {nome['weightCategoryFullName']}, realizado no dia {data_evento}, na cidade de {nome['local']}."
                            texto_eventos = texto_eventos.replace("Ã§", "ç")

                            lista_eventos.append(texto_eventos + "\n")
                except:

                    nome_do_cidadao = nome_do_cidadao+"sem resultados"

                    lista_eventos = ["sem resultados"]
                    print("atleta sem resultados")

                if not lista_eventos:

                    paragraph.text = 'O atleta não participou de eventos durante o período de recebimento da bolsa, até a presente data.'
                    paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
                    run = paragraph.runs[0]
                    font = run.font
                    font.name = 'Calibri'
                    font.size = Pt(11)
                    font.bold = False

                else:

                    paragraph.text = ''.join(map(str, lista_eventos)).encode('utf-8').decode('utf-8')
                    paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
                    run = paragraph.runs[0]
                    font = run.font
                    font.name = 'Calibri'
                    font.size = Pt(11)
                    font.bold = False


                print(lista_eventos)

            elif "Niterói/RJ" in paragraph.text:

                paragraph.text = f"Niterói/RJ, {current_date}".encode('utf-8').decode('utf-8')
                paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.RIGHT
                run = paragraph.runs[0]
                font = run.font
                font.name = 'Calibri'
                font.size = Pt(11)
                font.bold = False

            else:
                pass

        doc_path = fr"/bolsa_atleta_docs\new\{str(id_do_cidadao)}.docx"
        doc.save(doc_path)

        convert(doc_path)

        print("doc saved")


if __name__ == '__main__':

    current_date = datetime.now().strftime("%d de %B de %Y")

    current_month = datetime.now().month

    last_month = (datetime.now().replace(day=1) - timedelta(days=1)).month

    dados_ministerio = pd.read_excel(
        r'C:\Users\agata\PycharmProjects\RestApiBigmidia\tables\xlsx\pagamentos_pontuais_bolsa_atleta.xlsx')

    dados_ministerio['Fim do Pagamento'] = pd.to_datetime(dados_ministerio['Fim do Pagamento'], dayfirst=True)

    dados_do_mes = dados_ministerio[

        dados_ministerio['Fim do Pagamento'].dt.month == 4]  # ou current_month

    dados_do_mes["cpf_normal"] = dados_do_mes['CPF'].str.replace(r'\D', '', regex=True)

    create_declaration_prestacao(dados_do_mes)

